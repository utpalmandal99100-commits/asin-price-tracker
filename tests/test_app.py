from io import BytesIO

import pytest
from openpyxl import load_workbook

from app import create_app


class FakeAmazon:
    def __init__(self, on=True):
        self.on = on
        self.calls = []
        self.items = {
            "B0TEST1234": {
                "platform_id": "B0TEST1234", "title": "Maybelline New York Colossal Kajal, Black, 0.35g",
                "brand": "Maybelline", "variant": None, "eans": [], "model": None,
                "url": "https://www.amazon.in/dp/B0TEST1234", "mrp": 249.0, "price": 199.0,
                "currency": "INR", "status": "live", "note": None,
            }
        }

    def configured(self):
        return self.on

    def get_items(self, asins):
        self.calls.append(("get", list(asins)))
        return {a: self.items[a] for a in asins if a in self.items}

    def search(self, keywords, brand=None, count=10):
        self.calls.append(("search", keywords))
        return list(self.items.values())


@pytest.fixture
def make(tmp_path):
    def _make(on=True):
        fake = FakeAmazon(on)
        app = create_app(str(tmp_path / "t.db"), fake)
        return app.test_client(), fake
    return _make


def test_amazon_listing_fetched_live(make):
    c, fake = make()
    r = c.post("/api/products", json={"input": "https://www.amazon.in/dp/B0TEST1234"})
    p = r.get_json()["product"]
    l = p["listings"][0]
    assert (l["status"], l["price"], l["mrp"], l["source"]) == ("live", 199.0, 249.0, "Amazon Creators API")
    assert l["verified_at"]


def test_without_credentials_nothing_is_invented(make):
    c, _ = make(on=False)
    p = c.post("/api/products", json={"input": "B0TEST1234"}).get_json()["product"]
    l = p["listings"][0]
    assert l["status"] == "unverified" and l["price"] is None and l["verified_at"] is None


def test_manual_flow_suggested_match_and_lowest(make):
    c, fake = make()
    p = c.post("/api/products", json={"input": "https://www.nykaa.com/maybelline-colossal-kajal/p/14498"}).get_json()["product"]
    nykaa = p["listings"][0]
    assert nykaa["status"] == "unverified"

    p = c.post(f"/api/listings/{nykaa['id']}/manual", json={
        "title": "Maybelline New York Colossal Kajal Black 0.35 g", "brand": "Maybelline",
        "mrp": "249", "price": "189"}).get_json()
    amazon = next(l for l in p["listings"] if l["platform"] == "amazon")
    assert amazon["match_state"] == "suggested"
    # A suggested match never counts towards the lowest price.
    assert not amazon["counts_for_lowest"] and not any(l["is_lowest"] for l in p["listings"])

    c.post(f"/api/listings/{amazon['id']}/confirm")
    p = c.get("/api/products").get_json()[0]
    lowest = [l for l in p["listings"] if l["is_lowest"]]
    assert [l["platform"] for l in lowest] == ["nykaa"] and p["lowest_price"] == 189


def test_manual_rejects_bad_prices(make):
    c, _ = make(on=False)
    p = c.post("/api/products", json={"input": "11546402", "platform": "myntra"}).get_json()["product"]
    lid = p["listings"][0]["id"]
    assert c.post(f"/api/listings/{lid}/manual", json={"mrp": "500", "price": "900"}).status_code == 400
    assert c.post(f"/api/listings/{lid}/manual", json={"price": "abc"}).status_code == 400
    assert c.post(f"/api/listings/{lid}/manual", json={"price": ""}).status_code == 400


def test_ambiguous_id_asks_for_platform(make):
    c, _ = make()
    r = c.post("/api/products", json={"input": "11546402"})
    assert r.status_code == 409
    assert {x["platform"] for x in r.get_json()["candidates"]} == {"myntra", "nykaa"}


def test_refresh_is_rate_limited_and_batched(make):
    c, fake = make()
    c.post("/api/products", json={"input": "B0TEST1234"})
    fake.calls.clear()
    s = c.post("/api/refresh").get_json()["summary"]
    assert s["checked"] == 0 and fake.calls == []
    s = c.post("/api/refresh?force=1").get_json()["summary"]
    assert s["checked"] == 1 and fake.calls == [("get", ["B0TEST1234"])]


def test_api_failure_keeps_last_price_but_not_as_current(make):
    c, fake = make()
    c.post("/api/products", json={"input": "B0TEST1234"})

    def boom(asins):
        raise RuntimeError("throttled")
    fake.get_items = boom
    c.post("/api/refresh?force=1")
    l = c.get("/api/products").get_json()[0]["listings"][0]
    assert l["status"] == "error" and not l["counts_for_lowest"]


def test_excel_export(make):
    c, _ = make()
    c.post("/api/products", json={"input": "B0TEST1234"})
    r = c.get("/api/export.xlsx")
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.data))
    assert wb.sheetnames == ["Comparison", "Price History", "Read Me"]
    ws = wb["Comparison"]
    header = [c.value for c in ws[1]]
    assert header[1:6] == ["Product Name", "Platform", "Platform Product ID", "MRP", "Current Selling Price"]
    row = [c.value for c in ws[2]]
    assert row[2] == "Amazon India" and row[3] == "B0TEST1234" and row[5] == 199
    assert wb["Price History"].max_row == 2
