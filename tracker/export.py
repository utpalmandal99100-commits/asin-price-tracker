"""Excel export of everything tracked: comparison, price history and a key."""

from datetime import datetime, timedelta, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .platforms import PLATFORMS

IST = timezone(timedelta(hours=5, minutes=30))

STATUS_TEXT = {
    "live": "Verified (API)",
    "manual": "Verified (manual check)",
    "stale": "Stale manual check",
    "unavailable": "Unavailable",
    "unverified": "Not verified",
    "error": "Refresh failed",
}

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
LOWEST_FILL = PatternFill("solid", fgColor="D1FAE5")
WARN_FILL = PatternFill("solid", fgColor="FEF3C7")
MONEY = '"₹"#,##0.00'


def ist(ts):
    if not ts:
        return None
    return datetime.fromisoformat(ts).astimezone(IST).replace(tzinfo=None)


def _header(ws, headers, widths):
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def build_workbook(products, history_rows):
    wb = Workbook()

    ws = wb.active
    ws.title = "Comparison"
    headers = ["Product #", "Product Name", "Platform", "Platform Product ID", "MRP", "Current Selling Price",
               "Lowest Price", "Price Status", "Last Verified (IST)", "Data Source", "Match", "Match Score",
               "Note", "Product URL"]
    _header(ws, headers, [10, 48, 14, 22, 12, 14, 10, 22, 20, 22, 12, 11, 40, 50])
    for p in products:
        for l in p["listings"]:
            current = l["status"] in ("live", "manual")
            match = "Source" if l["role"] == "source" else ("Confirmed" if l["match_state"] == "confirmed" else "Suggested")
            ws.append([
                p["id"],
                l["title"] or p["name"] or "(name not verified)",
                PLATFORMS[l["platform"]],
                l["platform_id"],
                l["mrp"],
                l["price"] if current else None,
                "YES" if l["is_lowest"] else "",
                STATUS_TEXT.get(l["status"], l["status"]),
                ist(l["verified_at"]),
                l["source"] or "",
                match,
                l["match_score"],
                (f"Last verified price ₹{l['price']:,.2f}. " if not current and l["price"] is not None else "") + (l["note"] or ""),
                l["url"],
            ])
            r = ws.max_row
            ws.cell(r, 5).number_format = MONEY
            ws.cell(r, 6).number_format = MONEY
            ws.cell(r, 9).number_format = "dd-mmm-yyyy hh:mm"
            if l["is_lowest"]:
                for c in range(1, len(headers) + 1):
                    ws.cell(r, c).fill = LOWEST_FILL
                ws.cell(r, 6).font = Font(bold=True)
            elif not current:
                ws.cell(r, 8).fill = WARN_FILL
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

    hs = wb.create_sheet("Price History")
    _header(hs, ["Recorded (IST)", "Product #", "Product Name", "Platform", "Platform Product ID", "MRP",
                 "Selling Price", "Status", "Data Source"], [20, 10, 48, 14, 22, 12, 14, 22, 22])
    for h in history_rows:
        hs.append([ist(h["recorded_at"]), h["product_id"], h["title"], PLATFORMS[h["platform"]], h["platform_id"],
                   h["mrp"], h["price"], STATUS_TEXT.get(h["status"], h["status"]), h["source"]])
        r = hs.max_row
        hs.cell(r, 1).number_format = "dd-mmm-yyyy hh:mm"
        hs.cell(r, 6).number_format = MONEY
        hs.cell(r, 7).number_format = MONEY
    hs.auto_filter.ref = f"A1:I{max(hs.max_row, 1)}"

    ks = wb.create_sheet("Read Me")
    ks.column_dimensions["A"].width = 26
    ks.column_dimensions["B"].width = 100
    rows = [
        ("Exported (IST)", datetime.now(IST).strftime("%d-%b-%Y %H:%M")),
        ("", ""),
        ("Price Status", "Meaning"),
        ("Verified (API)", "Returned by the official Amazon Creators API at the Last Verified time."),
        ("Verified (manual check)", "Entered by you after checking the product page, at the Last Verified time."),
        ("Stale manual check", "A manual check older than the allowed age. Not used for Lowest Price."),
        ("Unavailable", "Out of stock or no offer at the last check."),
        ("Not verified", "No verified price yet. Selling price is left blank."),
        ("Refresh failed", "The last refresh failed. The last verified price is in the Note column only."),
        ("", ""),
        ("Lowest Price", "Lowest current selling price among verified, confirmed listings of the same product."),
        ("Match", "Source = the listing you added. Confirmed = you confirmed it is the same product. "
                  "Suggested = found automatically; not used for Lowest Price until you confirm it."),
    ]
    for a, b in rows:
        ks.append([a, b])
    for r in (1, 3):
        ks.cell(r, 1).font = Font(bold=True)
        ks.cell(r, 2).font = Font(bold=True)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
